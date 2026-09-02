import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE      = Path(__file__).parent
BRKGA     = BASE / "brkga/1_thread"
ROMULO    = BASE / "romulo"
OUT_MONO  = BASE / "romulo_mono.xlsx"
OUT_MULTI = BASE / "romulo_multi.xlsx"

METRICS          = ["sil", "db", "ch"]
METRIC_COL       = {"sil": "silhouette", "db": "davies-bouldin", "ch": "calinski-harabasz"}
HIGHER_IS_BETTER = {"sil": True, "db": False, "ch": True}

STAT_LABELS = ["Média", "Mínimo", "Q1", "Mediana", "Q3", "Máximo", "Vitórias"]

NC_COL = "number-clusters"

# ── data lake index ───────────────────────────────────────────────────────────

def build_lake_index() -> dict:
    """Returns {dataset_filename: lake_name}"""
    index = {}
    base = BASE.parent / "data-lakes" / "clustering"
    if not base.exists():
        base = Path("data-lakes") / "clustering"
    for csv in base.rglob("*.csv"):
        index[csv.name] = csv.parent.name
    return index

LAKE_INDEX = build_lake_index()

def get_lake(dataset_name: str) -> str:
    return LAKE_INDEX.get(dataset_name, "Unknown")

# ── data loading ──────────────────────────────────────────────────────────────

def read_run_files(folder: Path) -> list:
    if not folder.exists():
        return []
    files = sorted(folder.glob("*.csv"), key=lambda p: int(p.stem))
    return [pd.read_csv(f) for f in files]

def best_value(values, metric):
    return max(values) if HIGHER_IS_BETTER[metric] else min(values)

def is_best(v, best, metric):
    return (v >= best - 1e-12) if HIGHER_IS_BETTER[metric] else (v <= best + 1e-12)

# ── core stats (sil/db/ch) ────────────────────────────────────────────────────

def build_dataset_stats(brkga_runs, romulo_runs, metric):
    mcol = METRIC_COL[metric]

    all_names = set()
    for df in brkga_runs + romulo_runs:
        all_names |= set(df["dataset"].tolist())
    all_names = sorted(all_names)

    records = []
    for ds in all_names:
        def algo_data(runs):
            rows = []
            for df in runs:
                r = df[df["dataset"] == ds]
                if not r.empty:
                    rows.append((float(r[mcol].iloc[0]), float(r["time"].iloc[0])))
            return rows

        brkga_data  = algo_data(brkga_runs)
        romulo_data = algo_data(romulo_runs)

        if not brkga_data and not romulo_data:
            continue

        all_scores  = [s for s, _ in brkga_data + romulo_data]
        global_best = best_value(all_scores, metric)

        def summarise(data):
            if not data:
                return None, None
            scores = [s for s, _ in data]
            best_s = best_value(scores, metric)

            winning_times = [t for s, t in data if is_best(s, global_best, metric)]

            if winning_times:
                return best_s, min(winning_times)
            else:
                return best_s, 300.0

        brkga_score,  brkga_time  = summarise(brkga_data)
        romulo_score, romulo_time = summarise(romulo_data)

        brkga_wins  = brkga_score  is not None and is_best(brkga_score,  global_best, metric)
        romulo_wins = romulo_score is not None and is_best(romulo_score, global_best, metric)

        records.append({
            "lake":         get_lake(ds),
            "dataset":      ds,
            "brkga_score":  brkga_score,
            "brkga_time":   brkga_time,
            "romulo_score": romulo_score,
            "romulo_time":  romulo_time,
            "brkga_wins":   brkga_wins,
            "romulo_wins":  romulo_wins,
        })

    df = pd.DataFrame(records)
    if not df.empty:
        df = df.sort_values(["lake", "dataset"]).reset_index(drop=True)
    return df

# ── NC stats (number-clusters: brkga/sil vs romulo/original) ─────────────────

def build_nc_stats(brkga_sil_runs, romulo_original_runs):
    """
    Compare number-clusters from brkga/sil vs romulo/original.
    Lower is better. Returns a flat DataFrame sorted by dataset name.
    """
    all_names = set()
    for df in brkga_sil_runs + romulo_original_runs:
        if NC_COL in df.columns:
            all_names |= set(df["dataset"].tolist())
    all_names = sorted(all_names)

    records = []
    for ds in all_names:
        def algo_nc(runs):
            vals = []
            for df in runs:
                r = df[df["dataset"] == ds]
                if not r.empty and NC_COL in r.columns:
                    vals.append(float(r[NC_COL].iloc[0]))
            return vals

        brkga_vals  = algo_nc(brkga_sil_runs)
        romulo_vals = algo_nc(romulo_original_runs)

        if not brkga_vals and not romulo_vals:
            continue

        all_vals    = brkga_vals + romulo_vals
        global_best = min(all_vals)  # lower is better

        brkga_best  = min(brkga_vals)  if brkga_vals  else None
        romulo_best = min(romulo_vals) if romulo_vals else None

        brkga_wins  = brkga_best  is not None and brkga_best  <= global_best + 1e-12
        romulo_wins = romulo_best is not None and romulo_best <= global_best + 1e-12

        records.append({
            "dataset":     ds,
            "brkga_nc":    brkga_best,
            "romulo_nc":   romulo_best,
            "brkga_wins":  brkga_wins,
            "romulo_wins": romulo_wins,
        })

    return pd.DataFrame(records)

# ── Excel styling ─────────────────────────────────────────────────────────────

HEADER_FILL   = PatternFill("solid", start_color="2F5496", end_color="2F5496")
BRKGA_FILL    = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
ROMULO_FILL   = PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA")
AVG_FILL      = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
STAT_FILL     = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
WIN_FILL      = PatternFill("solid", start_color="FFE699", end_color="FFE699")
LAKE_AVG_FILL = PatternFill("solid", start_color="BDD7EE", end_color="BDD7EE")

THIN   = Side(style="thin",   color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="center")

def style(cell, bold=False, fill=None, align=CENTER, color="000000", size=10, fmt=None,
          border=BORDER):
    cell.font      = Font(name="Arial", bold=bold, color=color, size=size)
    cell.alignment = align
    cell.border    = border
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt

def stat_values(values):
    arr = np.array([v for v in values if v is not None], dtype=float)
    if len(arr) == 0:
        return {k: None for k in STAT_LABELS[:-1]}
    return {
        "Média":   float(np.mean(arr)),
        "Mínimo":  float(np.min(arr)),
        "Q1":      float(np.percentile(arr, 25)),
        "Mediana": float(np.median(arr)),
        "Q3":      float(np.percentile(arr, 75)),
        "Máximo":  float(np.max(arr)),
    }

# ── sheet writer (sil/db/ch) ──────────────────────────────────────────────────

def write_sheet(ws, df, metric, title):
    mcol = METRIC_COL[metric]

    ws.merge_cells("A1:F1")
    ws["A1"] = title
    style(ws["A1"], bold=True, fill=HEADER_FILL, color="FFFFFF", size=12)

    for col, val in enumerate(["", "", "BRKGA", "BRKGA", "Rómulo", "Rómulo"], 1):
        style(ws.cell(row=2, column=col, value=val),
              bold=True, fill=HEADER_FILL, color="FFFFFF")
    for col, val in enumerate(["Data Lake", "Dataset", mcol, "Tempo (s)", mcol, "Tempo (s)"], 1):
        style(ws.cell(row=3, column=col, value=val),
              bold=True, fill=HEADER_FILL, color="FFFFFF")

    current_row = 4
    lakes = df["lake"].unique()
    lake_ranges = {}

    for lake in lakes:
        group = df[df["lake"] == lake].reset_index(drop=True)
        first_row = current_row

        for _, row in group.iterrows():
            bw = row["brkga_wins"]
            rw = row["romulo_wins"]

            entries = [
                (row["lake"],         None,        None,     LEFT,   False),
                (row["dataset"],      None,        None,     LEFT,   False),
                (row["brkga_score"],  BRKGA_FILL,  "0.0000", CENTER, bw),
                (row["brkga_time"],   BRKGA_FILL,  "0.00",   CENTER, bw),
                (row["romulo_score"], ROMULO_FILL, "0.0000", CENTER, rw),
                (row["romulo_time"],  ROMULO_FILL, "0.00",   CENTER, rw),
            ]
            for col, (val, fl, fm, al, bold) in enumerate(entries, 1):
                style(ws.cell(row=current_row, column=col, value=val),
                      bold=bold, fill=fl, align=al, fmt=fm)
            current_row += 1

        last_row = current_row - 1
        lake_ranges[lake] = (first_row, last_row)

        avg_label = f"Média — {lake}"
        style(ws.cell(row=current_row, column=1, value=avg_label),
              bold=True, fill=LAKE_AVG_FILL, align=LEFT)
        ws.merge_cells(f"A{current_row}:B{current_row}")

        for col, fmt in [(3, "0.0000"), (4, "0.00"), (5, "0.0000"), (6, "0.00")]:
            cl = get_column_letter(col)
            c  = ws.cell(row=current_row, column=col,
                         value=f"=AVERAGE({cl}{first_row}:{cl}{last_row})")
            style(c, bold=True, fill=LAKE_AVG_FILL, fmt=fmt)

        current_row += 2  # blank separator

    # ── global summary ────────────────────────────────────────────────────────
    summary_start = current_row + 1

    bs = stat_values(df["brkga_score"].dropna().tolist())
    rs = stat_values(df["romulo_score"].dropna().tolist())
    bt = stat_values(df["brkga_time"].dropna().tolist())
    rt = stat_values(df["romulo_time"].dropna().tolist())

    brkga_wins_total  = int(df["brkga_wins"].sum())
    romulo_wins_total = int(df["romulo_wins"].sum())

    ws.merge_cells(f"A{summary_start}:F{summary_start}")
    style(ws.cell(row=summary_start, column=1, value="Estatísticas Globais"),
          bold=True, fill=HEADER_FILL, color="FFFFFF")

    sub = summary_start + 1
    for col, val in enumerate(["", "", "BRKGA", "BRKGA", "Rómulo", "Rómulo"], 1):
        style(ws.cell(row=sub, column=col, value=val),
              bold=True, fill=HEADER_FILL, color="FFFFFF")

    for offset, label in enumerate(STAT_LABELS):
        row_idx    = sub + 1 + offset
        is_win_row = (label == "Vitórias")
        is_avg_row = (label == "Média")
        row_fill   = WIN_FILL if is_win_row else (AVG_FILL if is_avg_row else STAT_FILL)

        style(ws.cell(row=row_idx, column=1, value=label),
              bold=True, fill=row_fill, align=LEFT)
        style(ws.cell(row=row_idx, column=2, value=""), fill=row_fill)

        if is_win_row:
            style(ws.cell(row=row_idx, column=3, value=brkga_wins_total),
                  bold=True, fill=WIN_FILL, fmt="0")
            style(ws.cell(row=row_idx, column=4, value=""), fill=WIN_FILL)
            style(ws.cell(row=row_idx, column=5, value=romulo_wins_total),
                  bold=True, fill=WIN_FILL, fmt="0")
            style(ws.cell(row=row_idx, column=6, value=""), fill=WIN_FILL)
        else:
            for col, (val, fl, fm) in enumerate([
                (bs[label], BRKGA_FILL,  "0.0000"),
                (bt[label], BRKGA_FILL,  "0.00"),
                (rs[label], ROMULO_FILL, "0.0000"),
                (rt[label], ROMULO_FILL, "0.00"),
            ], 3):
                style(ws.cell(row=row_idx, column=col, value=val),
                      bold=is_avg_row, fill=fl, fmt=fm)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 30
    for col in "CDEF":
        ws.column_dimensions[col].width = 18
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 18

# ── NC sheet writer ───────────────────────────────────────────────────────────

def write_nc_sheet(ws, df, title):
    """
    Flat layout: Dataset | BRKGA (number-clusters) | Rómulo (number-clusters)
    No lake grouping. Statistics block at the end. Lower = better.
    """
    # ── headers ───────────────────────────────────────────────────────────────
    ws.merge_cells("A1:C1")
    ws["A1"] = title
    style(ws["A1"], bold=True, fill=HEADER_FILL, color="FFFFFF", size=12)

    for col, val in enumerate(["Dataset", "BRKGA", "Rómulo"], 1):
        style(ws.cell(row=2, column=col, value=val),
              bold=True, fill=HEADER_FILL, color="FFFFFF")

    for col, val in enumerate(["", NC_COL, NC_COL], 1):
        style(ws.cell(row=3, column=col, value=val),
              bold=True, fill=HEADER_FILL, color="FFFFFF")

    # ── data rows ─────────────────────────────────────────────────────────────
    current_row = 4

    for _, row in df.iterrows():
        bw = row["brkga_wins"]
        rw = row["romulo_wins"]

        style(ws.cell(row=current_row, column=1, value=row["dataset"]),
              align=LEFT)
        style(ws.cell(row=current_row, column=2, value=row["brkga_nc"]),
              bold=bw, fill=BRKGA_FILL, fmt="0")
        style(ws.cell(row=current_row, column=3, value=row["romulo_nc"]),
              bold=rw, fill=ROMULO_FILL, fmt="0")
        current_row += 1

    # ── global summary ────────────────────────────────────────────────────────
    summary_start = current_row + 1

    bs = stat_values(df["brkga_nc"].dropna().tolist())
    rs = stat_values(df["romulo_nc"].dropna().tolist())

    brkga_wins_total  = int(df["brkga_wins"].sum())
    romulo_wins_total = int(df["romulo_wins"].sum())

    ws.merge_cells(f"A{summary_start}:C{summary_start}")
    style(ws.cell(row=summary_start, column=1, value="Estatísticas Globais"),
          bold=True, fill=HEADER_FILL, color="FFFFFF")

    sub = summary_start + 1
    for col, val in enumerate(["", "BRKGA", "Rómulo"], 1):
        style(ws.cell(row=sub, column=col, value=val),
              bold=True, fill=HEADER_FILL, color="FFFFFF")

    for offset, label in enumerate(STAT_LABELS):
        row_idx    = sub + 1 + offset
        is_win_row = (label == "Vitórias")
        is_avg_row = (label == "Média")
        row_fill   = WIN_FILL if is_win_row else (AVG_FILL if is_avg_row else STAT_FILL)

        style(ws.cell(row=row_idx, column=1, value=label),
              bold=True, fill=row_fill, align=LEFT)

        if is_win_row:
            style(ws.cell(row=row_idx, column=2, value=brkga_wins_total),
                  bold=True, fill=WIN_FILL, fmt="0")
            style(ws.cell(row=row_idx, column=3, value=romulo_wins_total),
                  bold=True, fill=WIN_FILL, fmt="0")
        else:
            style(ws.cell(row=row_idx, column=2, value=bs[label]),
                  bold=is_avg_row, fill=BRKGA_FILL, fmt="0.00")
            style(ws.cell(row=row_idx, column=3, value=rs[label]),
                  bold=is_avg_row, fill=ROMULO_FILL, fmt="0.00")

    # ── column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 18

# ── workbook builders ─────────────────────────────────────────────────────────

def build_mono():
    wb = Workbook()
    wb.remove(wb.active)

    for metric in METRICS:
        brkga_runs  = read_run_files(BRKGA  / metric)
        romulo_runs = read_run_files(ROMULO / metric)

        if not brkga_runs and not romulo_runs:
            print(f"  [mono/{metric}] Nenhum arquivo encontrado, pulando.")
            continue

        df = build_dataset_stats(brkga_runs, romulo_runs, metric)
        ws = wb.create_sheet(title=metric.upper())
        write_sheet(ws, df, metric, f"Rómulo Mono-objetivo — {METRIC_COL[metric]}")
        print(f"  [mono/{metric}] {len(df)} datasets processados.")

    wb.save(OUT_MONO)
    print(f"Salvo: {OUT_MONO}\n")


def build_multi():
    wb = Workbook()
    wb.remove(wb.active)

    romulo_original_runs = read_run_files(ROMULO / "original")

    for metric in METRICS:
        brkga_runs = read_run_files(BRKGA / metric)

        if not brkga_runs and not romulo_original_runs:
            print(f"  [multi/{metric}] Nenhum arquivo encontrado, pulando.")
            continue

        df = build_dataset_stats(brkga_runs, romulo_original_runs, metric)
        ws = wb.create_sheet(title=metric.upper())
        write_sheet(ws, df, metric, f"Rómulo Multi-objetivo — {METRIC_COL[metric]}")
        print(f"  [multi/{metric}] {len(df)} datasets processados.")

    # ── NC sheet: brkga/sil vs romulo/original ────────────────────────────────
    brkga_sil_runs = read_run_files(BRKGA / "sil")

    if brkga_sil_runs or romulo_original_runs:
        nc_df = build_nc_stats(brkga_sil_runs, romulo_original_runs)
        if not nc_df.empty:
            ws_nc = wb.create_sheet(title="NC")
            write_nc_sheet(ws_nc, nc_df,
                           "Número de Grupos — BRKGA/sil vs Rómulo/original")
            print(f"  [multi/nc] {len(nc_df)} datasets processados.")
        else:
            print("  [multi/nc] Nenhum dado encontrado para number-clusters.")
    else:
        print("  [multi/nc] Arquivos não encontrados, pulando.")

    wb.save(OUT_MULTI)
    print(f"Salvo: {OUT_MULTI}\n")


if __name__ == "__main__":
    print("=== Gerando romulo_mono.xlsx ===")
    build_mono()
    print("=== Gerando romulo_multi.xlsx ===")
    build_multi()
    print("Concluído.")